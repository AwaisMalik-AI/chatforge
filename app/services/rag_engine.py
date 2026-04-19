"""RAG: ingest, chunk, embed, ChromaDB search, prompt augmentation, groundedness."""

from __future__ import annotations

import csv
import hashlib
import io
import json
import logging
import re
from typing import Any

from bs4 import BeautifulSoup
from docx import Document as DocxDocument
from langchain_text_splitters import RecursiveCharacterTextSplitter
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.config import settings
from app.models.chatbot import (
    ChunkStrategy,
    DocumentSourceType,
    DocumentStatus,
    KnowledgeBase,
    KnowledgeChunk,
    KnowledgeDocument,
)

logger = logging.getLogger(__name__)

try:
    import fitz  # PyMuPDF
except ImportError:
    fitz = None  # type: ignore

try:
    import openpyxl
except ImportError:
    openpyxl = None  # type: ignore


class RAGEngine:
    def __init__(self) -> None:
        import chromadb

        raw = settings.VECTOR_DB_URL.strip()
        if raw.startswith("http://") or raw.startswith("https://"):
            from urllib.parse import urlparse

            u = urlparse(raw)
            port = u.port or (443 if u.scheme == "https" else 80)
            ssl = u.scheme == "https"
            self._client = chromadb.HttpClient(host=u.hostname or "localhost", port=port, ssl=ssl)
        else:
            self._client = chromadb.PersistentClient(path=raw)
        self._st_model = None

    def _embedder(self):
        if self._st_model is None:
            from sentence_transformers import SentenceTransformer

            self._st_model = SentenceTransformer(settings.EMBEDDING_MODEL)
        return self._st_model

    def _collection_name(self, kb_id: int) -> str:
        return f"kb_{kb_id}"

    def _get_collection(self, kb_id: int):
        return self._client.get_or_create_collection(
            name=self._collection_name(kb_id),
            metadata={"hnsw:space": "cosine"},
        )

    def _splitter_for_kb(self, kb: KnowledgeBase):
        overlap = kb.chunk_overlap
        size = kb.chunk_size
        if kb.chunk_strategy == ChunkStrategy.fixed:
            return RecursiveCharacterTextSplitter(
                chunk_size=size,
                chunk_overlap=overlap,
                separators=["\n\n", "\n", ". ", " ", ""],
            )
        if kb.chunk_strategy == ChunkStrategy.sentence:
            return RecursiveCharacterTextSplitter(
                chunk_size=size,
                chunk_overlap=overlap,
                separators=[". ", ".\n", "\n\n", "\n", " ", ""],
            )
        return RecursiveCharacterTextSplitter(
            chunk_size=size * 4,
            chunk_overlap=overlap,
            separators=["\n\n", "\n"],
        )

    def _extract_pdf(self, file_bytes: bytes) -> tuple[str, dict[str, Any]]:
        meta: dict[str, Any] = {}
        if not fitz:
            return file_bytes.decode("utf-8", errors="replace"), meta
        doc = fitz.open(stream=file_bytes, filetype="pdf")
        try:
            parts = [page.get_text() for page in doc]
            m = doc.metadata
            if m:
                if m.get("title"):
                    meta["title"] = str(m["title"]).strip()
                if m.get("author"):
                    meta["author"] = str(m["author"]).strip()
        finally:
            doc.close()
        return "\n".join(parts), meta

    def _extract_docx(self, file_bytes: bytes) -> tuple[str, dict[str, Any]]:
        meta: dict[str, Any] = {}
        d = DocxDocument(io.BytesIO(file_bytes))
        text = "\n".join(p.text for p in d.paragraphs)
        core = d.core_properties
        if core:
            if core.title:
                meta["title"] = str(core.title).strip()
            if core.subject:
                meta["subject"] = str(core.subject).strip()
            if core.keywords:
                meta["keywords"] = str(core.keywords).strip()
            if core.version:
                meta["version"] = str(core.version).strip()
            if core.last_modified_by:
                meta["last_modified_by"] = str(core.last_modified_by).strip()
        return text, meta

    def _extract_html(self, file_bytes: bytes) -> tuple[str, dict[str, Any]]:
        raw = file_bytes.decode("utf-8", errors="replace")
        soup = BeautifulSoup(raw, "html.parser")
        text = soup.get_text(separator="\n", strip=True)
        meta: dict[str, Any] = {}
        for tag in soup.find_all("meta"):
            name = (tag.get("name") or tag.get("property") or "").lower()
            content = tag.get("content")
            if not content or not name:
                continue
            if name in ("department", "dc.creator", "article:tag"):
                meta.setdefault("department", str(content).strip())
            if name in ("version", "doc-version", "document-version"):
                meta["version"] = str(content).strip()
            if name in ("sensitivity", "classification", "security", "data-classification"):
                meta["sensitivity"] = str(content).strip()
        return text, meta

    def _extract_spreadsheet(self, file_bytes: bytes, filename: str) -> tuple[str, dict[str, Any]]:
        name = (filename or "").lower()
        meta: dict[str, Any] = {}
        if name.endswith(".csv"):
            s = file_bytes.decode("utf-8-sig", errors="replace")
            reader = csv.reader(io.StringIO(s))
            rows = list(reader)
            if rows:
                meta["columns"] = ",".join(str(c) for c in rows[0])
            lines = [",".join(row) for row in rows[:5000]]
            return "\n".join(lines), meta

        if name.endswith(".xlsx") and openpyxl:
            wb = openpyxl.load_workbook(io.BytesIO(file_bytes), read_only=True, data_only=True)
            try:
                props = wb.properties
                ver = getattr(props, "version", None) if props else None
                if ver:
                    meta["version"] = str(ver)
                parts: list[str] = []
                for sheet in wb.worksheets:
                    parts.append(f"## Sheet: {sheet.title}")
                    for i, row in enumerate(sheet.iter_rows(values_only=True)):
                        if i >= 2000:
                            break
                        cells = [str(c) if c is not None else "" for c in row]
                        parts.append("\t".join(cells))
                return "\n".join(parts), meta
            finally:
                wb.close()

        if name.endswith(".xlsx") and not openpyxl:
            logger.warning("openpyxl not installed; skipping structured XLSX parse")

        return file_bytes.decode("utf-8", errors="replace"), meta

    def _yaml_front_matter(self, text: str) -> tuple[str, dict[str, Any]]:
        meta: dict[str, Any] = {}
        if not text.startswith("---"):
            return text, meta
        end = text.find("\n---", 3)
        if end == -1:
            return text, meta
        block = text[3:end].strip()
        body = text[end + 4 :].lstrip("\n")
        for line in block.splitlines():
            if ":" not in line:
                continue
            k, v = line.split(":", 1)
            key = k.strip().lower().replace(" ", "_")
            val = v.strip().strip('"').strip("'")
            if key in ("department", "version", "sensitivity", "classification"):
                if key == "classification":
                    meta["sensitivity"] = val
                else:
                    meta[key] = val
        return body, meta

    def _chroma_metadata(self, meta: dict[str, Any]) -> dict[str, Any]:
        """Chroma only accepts str, int, float, bool; coerce the rest."""
        out: dict[str, Any] = {}
        for k, v in meta.items():
            if v is None:
                continue
            if isinstance(v, (str, int, float, bool)):
                out[k] = v
            else:
                out[k] = json.dumps(v, default=str)[:2000]
        return out

    def _merge_enterprise_metadata(self, *sources: dict[str, Any]) -> dict[str, Any]:
        out: dict[str, Any] = {}
        for s in sources:
            if not s:
                continue
            for key in ("department", "version", "sensitivity", "title", "author"):
                if s.get(key) and key not in out:
                    out[key] = s[key]
        return out

    def extract_text(
        self,
        file_bytes: bytes,
        filename: str,
        source_type: DocumentSourceType,
    ) -> tuple[str, dict[str, Any]]:
        name = (filename or "doc").lower()
        if source_type == DocumentSourceType.text:
            raw = file_bytes.decode("utf-8", errors="replace")
            body, fm = self._yaml_front_matter(raw)
            return body, self._merge_enterprise_metadata(fm, {})

        if name.endswith(".pdf"):
            text, m = self._extract_pdf(file_bytes)
            return text, self._merge_enterprise_metadata(m, {})

        if name.endswith(".docx"):
            text, m = self._extract_docx(file_bytes)
            body, fm = self._yaml_front_matter(text)
            return body, self._merge_enterprise_metadata(m, fm)

        if name.endswith((".html", ".htm")):
            text, m = self._extract_html(file_bytes)
            body, fm = self._yaml_front_matter(text)
            return body, self._merge_enterprise_metadata(m, fm)

        if name.endswith((".xlsx", ".csv")):
            text, m = self._extract_spreadsheet(file_bytes, filename)
            body, fm = self._yaml_front_matter(text)
            return body, self._merge_enterprise_metadata(m, fm)

        if name.endswith(".txt"):
            raw = file_bytes.decode("utf-8", errors="replace")
            body, fm = self._yaml_front_matter(raw)
            return body, fm

        raw = file_bytes.decode("utf-8", errors="replace")
        body, fm = self._yaml_front_matter(raw)
        return body, fm

    async def ingest_document(
        self,
        db: AsyncSession,
        kb_id: int,
        file_bytes: bytes,
        filename: str,
        source_type: DocumentSourceType,
        title: str | None = None,
    ) -> KnowledgeDocument:
        kb = await db.get(KnowledgeBase, kb_id)
        if not kb:
            raise ValueError("Knowledge base not found")
        text, enterprise_meta = self.extract_text(file_bytes, filename, source_type)
        h = hashlib.sha256(text.encode()).hexdigest()
        doc = KnowledgeDocument(
            kb_id=kb_id,
            title=title or filename,
            source_type=source_type,
            source_path=filename,
            content_hash=h,
            status=DocumentStatus.processing,
            enterprise_metadata=enterprise_meta or None,
        )
        db.add(doc)
        await db.flush()

        splitter = self._splitter_for_kb(kb)
        chunks_text = splitter.split_text(text)
        model = self._embedder()
        embeddings = model.encode(chunks_text, show_progress_bar=False).tolist()
        coll = self._get_collection(kb_id)
        ids = [f"d{doc.id}_c{i}" for i in range(len(chunks_text))]
        base_chunk_meta = self._chroma_metadata({"title": doc.title, **(enterprise_meta or {})})
        coll.add(
            ids=ids,
            embeddings=embeddings,
            documents=chunks_text,
            metadatas=[
                {"document_id": doc.id, "chunk_index": i, **base_chunk_meta}
                for i in range(len(chunks_text))
            ],
        )
        for i, (ct, emb) in enumerate(zip(chunks_text, embeddings, strict=False)):
            ch = KnowledgeChunk(
                document_id=doc.id,
                content=ct,
                embedding_vector=json.dumps(emb),
                chunk_index=i,
                chunk_metadata={"title": doc.title, **(enterprise_meta or {})},
                token_count=max(1, len(ct) // 4),
            )
            db.add(ch)
        doc.status = DocumentStatus.indexed
        doc.chunk_count = len(chunks_text)
        kb.total_documents = (kb.total_documents or 0) + 1
        kb.total_chunks = (kb.total_chunks or 0) + len(chunks_text)
        await db.flush()
        return doc

    async def search(
        self,
        db: AsyncSession,
        kb_id: int,
        query: str,
        top_k: int = 5,
        score_threshold: float = 0.0,
        metadata_filter: dict[str, Any] | None = None,
    ) -> list[dict[str, Any]]:
        kb = await db.get(KnowledgeBase, kb_id)
        if not kb or not kb.is_active:
            return []
        model = self._embedder()
        q_emb = model.encode([query], show_progress_bar=False).tolist()[0]
        coll = self._get_collection(kb_id)
        where = metadata_filter
        res = coll.query(
            query_embeddings=[q_emb],
            n_results=top_k,
            where=where,
            include=["documents", "metadatas", "distances"],
        )
        hits: list[dict[str, Any]] = []
        if not res["documents"] or not res["documents"][0]:
            return hits
        for doc_txt, meta, dist in zip(
            res["documents"][0],
            res["metadatas"][0] or [],
            res["distances"][0],
            strict=False,
        ):
            score = 1.0 - float(dist) if dist is not None else 0.0
            if score < score_threshold:
                continue
            hits.append(
                {
                    "content": doc_txt,
                    "score": score,
                    "metadata": meta,
                    "chunk_id": meta.get("chunk_id") if isinstance(meta, dict) else None,
                }
            )
        return hits

    def augment_prompt(self, query: str, search_results: list[dict[str, Any]]) -> str:
        if not search_results:
            return query
        parts = []
        for i, r in enumerate(search_results, 1):
            src = r.get("metadata") or {}
            title = src.get("title", "source")
            parts.append(f"[{i}] ({title}): {r.get('content', '')}")
        ctx = "\n\n".join(parts)
        return (
            "Use the following retrieved context to answer. Cite sources as [n] when applicable.\n\n"
            f"Context:\n{ctx}\n\nUser question: {query}"
        )

    def evaluate_groundedness(self, response: str, sources: list[str]) -> float:
        if not response.strip():
            return 0.0
        if not sources:
            return 0.3
        resp_words = set(re.findall(r"\w+", response.lower()))
        joined = " ".join(sources).lower()
        overlap = sum(1 for w in resp_words if len(w) > 3 and w in joined)
        return min(1.0, overlap / max(8, len(resp_words) * 0.15))


class _RAGProxy:
    _inst: RAGEngine | None = None

    def _get(self) -> RAGEngine:
        if self._inst is None:
            self._inst = RAGEngine()
        return self._inst

    def __getattr__(self, name: str):
        return getattr(self._get(), name)


rag_engine = _RAGProxy()
